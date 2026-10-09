"""
Engine tests.  Run from the project folder:   pytest -q

The first test runs the notebook v11.7 self-test inside the engine (119 checks, including the heat balance, the heat
audit and the slow tornado and break-even checks), so a pass here means the deployed engine behaves exactly like the notebook.  The rest cover what
the dashboard adds: the Config / session mechanism, thread safety, the loader and template, the run bundle
(0-100 % sinter curve, sinter decision, oxide sources), the studio functions and the in-memory Excel export.
"""
import io
import os
import sys
import threading

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from mbf import optimiser as o
from mbf.optimiser import Config


def cfg(**kw):
    c = Config()
    for k, v in kw.items():
        assert hasattr(c, k), k
        setattr(c, k, v)
    return c


def base():
    return o.demo_df()


# ---------------------------------------------------------------- the notebook's own self-test
def test_notebook_v117_self_test_passes_in_full():
    with o.session(Config()):
        assert o.run_self_tests(verbose=False, full=True)


def test_quick_self_test_passes():
    assert all(ok for _, ok, _ in o.quick_self_test())


def test_demo_regression_anchor():
    """Pins the demo answer so any change to the rules or solver is noticed. The v11.7 heat balance in check mode must
    leave the v11.5 answer unchanged."""
    st, blend, cost, ach, diag = o.solve(base(), Config())
    assert st == "Optimal"
    assert abs(cost - 20033.29) < 0.05 and abs(ach["Sinter_share_pct"] - 76.7) < 0.05
    assert abs(ach["Fuel_supplied"] - ach["Fuel_rule"]) < 0.05
    assert any(x.startswith("SINTER DECISION") for x in diag)


# ---------------------------------------------------------------- Config, sessions, threads
def test_config_defaults_are_the_plant_values():
    c = Config()
    assert c.problems() == []
    assert (c.sinter_min, c.sinter_max, c.pci_fixed_kgthm, c.nut_coke_kgthm) == (0.50, 0.80, 120.0, 40.0)
    assert (c.basicity_min, c.basicity_max, c.mgo_min_pct, c.mgo_max_pct, c.al2o3_min_pct, c.al2o3_max_pct) == (0.99, 1.01, 7.0, 8.0, 17.0, 18.5)
    assert (c.fuel_per_pct_sinter, c.fuel_per_kg_raw_flux, c.fuel_per_kg_slag, c.ks_fixed) == (1.0, 0.30, 0.18, 25.0)
    assert all(v["base"] == 545.0 for v in c.furnace_profiles.values())
    assert set(c.fuel_terms) == set(o.RULE_SWITCHES) and all(c.fuel_terms.values())
    assert not hasattr(c, "hbt_actual") and not hasattr(c, "fe_form") and not hasattr(c, "fuel_offset_kgthm")


def test_config_problems_are_reported_in_plain_language():
    assert any("minimum is above the maximum" in p for p in cfg(basicity_min=1.2, basicity_max=1.0).problems())
    assert any("guard rails" in p for p in cfg(sinter_min=0.8, sinter_max=0.7).problems())
    assert any("between 0 and 100" in p for p in cfg(sinter_manual_on=True, sinter_manual_pct=101.0).problems())
    assert cfg(sinter_manual_on=True, sinter_manual_pct=100.0).problems() == []          # 100 % sinter can be pinned
    assert any("Stock balance" in p for p in cfg(stock_balance=1.5).problems())


def test_session_restores_module_defaults_even_after_an_error():
    before = (o.SINTER_MIN, o.PCI_FIXED_KGTHM, dict(o.STOCK), dict(o.FUEL_TERMS), o.REQUIRED_FE_KGTHM, o.KS_FIXED)
    with pytest.raises(RuntimeError):
        with o.session(cfg(sinter_min=0.55, pci_fixed_kgthm=99.0, stock_balance=0.0, fe_required_per_100kg=95.0, ks_fixed=30.0)):
            assert o.SINTER_MIN == 0.55 and o.REQUIRED_FE_KGTHM == 950.0 and o.KS_FIXED == 30.0
            raise RuntimeError("boom")
    assert (o.SINTER_MIN, o.PCI_FIXED_KGTHM, dict(o.STOCK), dict(o.FUEL_TERMS), o.REQUIRED_FE_KGTHM, o.KS_FIXED) == before


def test_new_config_after_a_session_still_has_plant_defaults():
    c = cfg()
    c.material_rules["ore_fe"]["coef"] = 99.0
    c.fuel_terms["slag"] = False
    with o.session(c):
        pass
    assert Config().material_rules["ore_fe"]["coef"] == 3.0 and Config().fuel_terms["slag"] is True
    assert o.MATERIAL_RULES["ore_fe"]["coef"] == 3.0 and o.FUEL_TERMS["slag"] is True


def test_settings_change_the_answer_as_expected():
    st, _, cost_a, a, _ = o.solve(base(), cfg(pci_fixed_kgthm=120.0))
    st2, _, cost_b, b, _ = o.solve(base(), cfg(pci_fixed_kgthm=100.0))
    assert st == st2 == "Optimal" and abs(b["PCI_kg"] - 100.0) < 1e-6 and b["Coke_kg"] > a["Coke_kg"]
    s3, _, _, c3, _ = o.solve(base(), cfg(fe_required_per_100kg=95.0))
    assert s3 == "Optimal" and abs(c3["Fe_burden_kg"] - 950.0) < 1e-3
    c = cfg(furnace_name="MBF-2"); c.furnace_profiles["MBF-2"]["base"] = 540.0
    s4, _, _, c4, _ = o.solve(base(), c)
    assert s4 == "Optimal" and c4["Fuel_terms"]["base"] == 540.0
    s5, _, _, c5, _ = o.solve(base(), cfg(ks_fixed=50.0))
    assert s5 == "Optimal" and c5["S_HM_pct"] < a["S_HM_pct"]


def test_rule_switch_in_config_removes_that_term_and_is_named():
    c = cfg(); c.fuel_terms["other_moisture"] = False; c.fuel_terms["raw_flux"] = False
    st, _, _, a, diag = o.solve(base(), c)
    assert st == "Optimal" and a["Fuel_terms"]["other_moisture"] == 0 and a["Fuel_terms"]["raw_flux"] == 0
    assert any("SWITCHED OFF" in x and "Raw flux" in x for x in diag)
    assert c.rules_off() == ["Raw flux", "Sinter/flux moisture"]


def test_concurrent_sessions_do_not_leak_settings():
    """Three threads with different settings must each get exactly what they'd get alone."""
    cases = [cfg(pci_fixed_kgthm=100.0, sinter_manual_on=True, sinter_manual_pct=68.0),
             cfg(pci_fixed_kgthm=140.0, fe_required_per_100kg=95.5, furnace_name="MBF-2"),
             cfg(al2o3_max_pct=18.2, price_basis="wet")]
    alone = [o.solve(base(), c, explain=False) for c in cases]
    got = {}

    def work(i, c):
        for _ in range(3):
            got[i] = o.solve(base(), c, explain=False)

    ts = [threading.Thread(target=work, args=(i, c)) for i, c in enumerate(cases)]
    [t.start() for t in ts]; [t.join() for t in ts]
    for i in range(3):
        assert got[i][0] == alone[i][0]
        assert abs(got[i][2] - alone[i][2]) < 1e-9, i
        assert abs(got[i][3]["PCI_kg"] - alone[i][3]["PCI_kg"]) < 1e-9


# ---------------------------------------------------------------- loader, template, legacy tables
def test_template_round_trip_matches_demo_table():
    df, notes = o.load_master(o.template_bytes())
    d0 = base()
    assert list(df.columns) == o.COLUMNS[1:] and list(df.index) == list(d0.index)
    assert np.allclose(df[o.NUM_COLS].values, d0[o.NUM_COLS].values)
    assert "Ash %" not in pd.read_excel(io.BytesIO(o.template_bytes())).columns


def test_loader_group_spelling_legacy_columns_and_notes():
    t = base().reset_index()
    t["Availability"] = np.where(t["Available"], "ON", "OFF")
    t = t.drop(columns=["Available"])
    t.loc[0, "Group"] = "iron ore"
    t["Ash %"] = 12.0; t["CSR"] = 64.0
    t.loc[2, "Availability"] = None                      # Sinter1 (already ON): blank must read as ON
    bio = io.BytesIO(); t.to_excel(bio, index=False)
    df, notes = o.load_master(bio.getvalue())
    assert df.loc["Ore1", "Group"] == "Iron_ore"
    assert any("read as 'Iron_ore'" in n for n in notes) and any("ignored" in n for n in notes) and any("blank" in n for n in notes)
    assert "Ash_Pct" not in df.columns and "CSR" not in df.columns


def test_old_table_with_ash_and_csr_columns_still_solves():
    d = base(); d["Ash_Pct"] = np.nan; d["CSR"] = np.nan
    assert o.solve(d, Config())[0] == "Optimal"
    assert list(o.ensure_columns(d).columns) == o.COLUMNS[1:]


def test_input_checks_flag_an_expensive_flux():
    d = base(); d.loc["Quartzite", "Price_Rs_t"] = 11000.0
    diag = o.solve(d, Config())[4]
    assert any(x.startswith("CHECK: Quartzite") for x in diag)


# ---------------------------------------------------------------- the run bundle
def test_bundle_has_curve_decision_and_oxide_sources():
    res = o.solve(base(), Config())
    b = o.make_bundle(base(), res, Config())
    cv = b["curve"]
    assert len(cv) == 41 and cv["Sinter %"].iloc[0] == 0.0 and cv["Sinter %"].iloc[-1] == 100.0
    assert cv.attrs["hi"] is not None and 50 < cv.attrs["hi"] < 100 and "blocked by" in cv.attrs["note"]
    assert b["choice"][0] == pytest.approx(res[3]["Sinter_share_pct"]) and b["choice"][1] == pytest.approx(res[2])
    assert b["decision"].startswith("SINTER DECISION")
    co = b["compact_oxides"]
    assert list(co.columns[2:]) == ["CaO kg", "CaO % of total", "SiO2 kg", "SiO2 % of total", "Al2O3 kg", "Al2O3 % of total"]
    tot = co[co["Material"] == "TOTAL CHARGED"].iloc[0]
    assert abs(float(tot["CaO kg"]) - res[3]["CaO_kg"]) < 0.01
    assert isinstance(b["cfg"], Config) and b["rules_off"] == []


def test_infeasible_run_still_gets_a_curve():
    c = cfg(sinter_manual_on=True, sinter_manual_pct=95.0)
    res = o.solve(base(), c)
    assert res[0] == "Infeasible"
    b = o.make_bundle(base(), res, c)
    assert b["curve"] is not None and b["choice"] is None


def test_curve_and_sweep_progress_callbacks():
    calls = []
    o.curve(base(), Config(), step=25.0, progress=lambda i, n: calls.append((i, n)))
    assert calls[-1] == (5, 5)
    calls.clear()
    out, oxl = o.sweep(base(), Config(), 60, 70, 5, progress=lambda i, n: calls.append((i, n)))
    assert calls[-1] == (3, 3) and len(out) == 4 and len(oxl) > 0


# ---------------------------------------------------------------- studio functions
def test_studio_functions_run_and_restore_every_setting():
    before = (o.REQUIRED_FE_KGTHM, o.AL2O3_MAX_PCT, o.MGO_MAX_PCT, o.BASICITY_MIN, o.BASICITY_MAX, o.FUEL_PER_KG_SLAG,
              {k: v["coef"] for k, v in o.MATERIAL_RULES.items()}, {k: dict(v) for k, v in o.FURNACE_PROFILES.items()})
    ps = o.price_sens(base(), Config(), "Group: Sinter", 10)
    assert (ps["Status"] == "Optimal").all() and ps["Cost Rs/tHM"].is_monotonic_increasing
    an = o.sensitivity(base(), Config(), "Iron_ore", "Al2O3", 0.5)
    assert len(an) == 5
    tor = o.tornado_run(base(), Config(), 10)
    assert tor["Swing Rs"].is_monotonic_decreasing and (tor["Status"] == "Optimal").all()
    be = o.breakeven(base(), Config())
    assert be.iloc[0]["Case"] == "Entered sinter price"
    after = (o.REQUIRED_FE_KGTHM, o.AL2O3_MAX_PCT, o.MGO_MAX_PCT, o.BASICITY_MIN, o.BASICITY_MAX, o.FUEL_PER_KG_SLAG,
             {k: v["coef"] for k, v in o.MATERIAL_RULES.items()}, {k: dict(v) for k, v in o.FURNACE_PROFILES.items()})
    assert before == after


def test_breakeven_needs_the_free_choice():
    with pytest.raises(ValueError):
        o.breakeven(base(), cfg(sinter_manual_on=True, sinter_manual_pct=70.0))


# ---------------------------------------------------------------- Excel export
def test_excel_export_has_every_sheet_and_the_right_numbers():
    from openpyxl import load_workbook
    import datetime as dt
    res = o.solve(base(), Config())
    b = o.make_bundle(base(), res, Config())
    now = dt.datetime.now()
    out, oxl = o.sweep(base(), Config(), 60, 80, 10)
    ana = {"sweep": {"df": out, "time": now, "fingerprint": b["fingerprint"], "title": "Sinter sweep"},
           "sweep_oxides": {"df": oxl, "time": now, "fingerprint": b["fingerprint"], "title": "Oxides"},
           "price": {"df": o.price_sens(base(), Config(), "Sinter1", 10), "time": now, "fingerprint": b["fingerprint"], "title": "Price"}}
    wb = load_workbook(io.BytesIO(o.export_bytes(b, ana)))
    assert {"Summary", "Sinter Curve 0-100%", "Inputs", "Optimised Burden", "Slag & Chemistry", "Fuel Rate", "Fe Impact",
            "Moisture", "Sinter Sweep", "Oxides by Sinter %", "Price Sensitivity", "Run Settings"} <= set(wb.sheetnames)
    assert len(wb["Sinter Curve 0-100%"]._charts) == 2
    ws = wb["Summary"]
    cell = next(((r, c) for r in range(1, 30) for c in range(1, 8) if ws.cell(r, c).value == "Total raw-material cost (Rs/tHM)"))
    assert abs(ws.cell(cell[0] + 1, cell[1]).value - res[2]) < 1e-6


def test_export_uses_the_settings_frozen_in_the_bundle():
    from openpyxl import load_workbook
    c = cfg(pci_fixed_kgthm=100.0)
    b = o.make_bundle(base(), o.solve(base(), c), c, with_curve=False)
    wb = load_workbook(io.BytesIO(o.export_bytes(b)))
    ws = wb["Run Settings"]
    vals = {ws.cell(r, 2).value: ws.cell(r, 3).value for r in range(1, ws.max_row + 1)}
    assert vals.get("PCI (kg/tHM)") == 100.0
    assert o.PCI_FIXED_KGTHM == 120.0                    # module defaults untouched afterwards


def test_failed_run_still_exports_inputs_and_reason():
    from openpyxl import load_workbook
    d = base(); d.loc["Coke3", "Available"] = False
    res = o.solve(d, Config())
    wb = load_workbook(io.BytesIO(o.export_bytes(o.make_bundle(d, res, Config()))))
    assert {"Summary", "Inputs", "Run Settings"} <= set(wb.sheetnames) and "Optimised Burden" not in wb.sheetnames


# ---------------------------------------------------------------- v11.7: heat balance, heat audit, every input in Config
def test_config_carries_the_heat_settings_and_model_constants():
    c = Config()
    assert c.heat == o.HEAT_DEFAULT and c.heat is not o.HEAT_DEFAULT and c.heat["mode"] == "check"
    assert c.ref_rates == {"ore": 600.0, "sinter": 1115.0, "coke": 402.0}
    assert (c.raw_flux_min_cao_mgo, c.dr_degree_floor, c.mn_reduction_eff) == (15.0, 0.40, 0.85)
    assert (c.mgo_al2o3_guide_lo, c.mgo_al2o3_guide_hi, c.fe_c_target, c.fe_c_tol) == (0.40, 0.55, 2.0, 0.1)
    assert set(o.HEAT_FIELDS) == set(o.HEAT_NUMERIC_KEYS)
    bad = cfg(); bad.heat["mode"] = "hot"; bad.heat["drr"] = 1.4; bad.ref_rates["ore"] = 0.0
    p = " ".join(bad.problems())
    assert "mode" in p and "direct reduction" in p and "Reference charge rates" in p


def test_heat_settings_never_leak_between_sessions():
    before = (o.copy.deepcopy(o.HEAT), dict(o.REF_RATES), o.MGO_AL2O3_GUIDE, o.RAW_FLUX_MIN_CAO_MGO)
    c = cfg(mgo_al2o3_guide_lo=0.3, raw_flux_min_cao_mgo=20.0)
    c.heat.update(mode="floor", loss_MJ_tHM=900.0, calibrated=True); c.ref_rates["ore"] = 650.0
    with o.session(c):
        assert o.HEAT["mode"] == "floor" and o.HEAT["loss_MJ_tHM"] == 900.0 and o.REF_RATES["ore"] == 650.0
        assert o.MGO_AL2O3_GUIDE == (0.3, 0.55) and o.RAW_FLUX_MIN_CAO_MGO == 20.0
    assert (o.HEAT, dict(o.REF_RATES), o.MGO_AL2O3_GUIDE, o.RAW_FLUX_MIN_CAO_MGO) == before
    assert Config().heat["loss_MJ_tHM"] == 500.0


def test_check_mode_reports_and_floor_mode_raises_coke():
    st, _, cost, a, diag = o.solve(base(), Config())
    h = a["Heat"]
    assert any(x.startswith("HEAT CHECK") for x in diag) and h["mode"] == "check" and not h["calibrated"]
    c = cfg(); c.heat["mode"] = "floor"
    st2, _, cost2, a2, diag2 = o.solve(base(), c)
    assert st2 == "Optimal" and a2["Heat"]["C_surplus_kg"] >= -1e-6
    if h["C_surplus_kg"] < 0:                           # placeholders leave a deficit on the demo data
        assert a2["Coke_kg"] > a["Coke_kg"] and cost2 > cost and any(x.startswith("HEAT BALANCE FLOOR") for x in diag2)


def test_calibrate_losses_closes_the_run_and_changes_nothing_else():
    st, blend, cost, a, _ = o.solve(base(), Config())
    new = o.calibrate_losses(a, Config())
    assert new["calibrated"] and new["loss_MJ_tHM"] != 500.0 and o.HEAT["loss_MJ_tHM"] == 500.0
    c = cfg(); c.heat = new
    st2, _, cost2, a2, _ = o.solve(base(), c)
    assert abs(a2["Heat"]["C_surplus_kg"]) < 1e-6 and abs(cost2 - cost) < 1e-6      # check mode: same burden, balance closed


def test_heat_audit_wrappers_round_trip_and_apply_per_session():
    from openpyxl import load_workbook
    tpl = o.audit_template_bytes(example=True)
    rec = o.audit_read(tpl, "example.xlsx")
    assert rec["name"] == "example.xlsx" and len(rec["cooling"]) == 3
    with pytest.raises(ValueError, match="missing required"):
        o.audit_read(o.audit_template_bytes(example=False))
    recs = o.audit_examples()
    A = o.audit_run_cfg(recs, Config())
    assert A["fits"]["drr"]["applied"] and abs(A["fits"]["drr"]["slope"] + 0.004) < 4e-4
    assert A["fit_text"] and len(A["sensitivity"]) == 4
    heat, note = o.audit_apply_cfg(A, Config(), "measured")
    assert heat["calibrated"] and "audited on 4" in note and o.HEAT == o.HEAT_DEFAULT
    assert {"Months", "Fits", "Flags", "Heat terms"} <= set(load_workbook(io.BytesIO(o.audit_export_bytes(A))).sheetnames)
    c = cfg(); c.heat = heat
    st, _, _, a, _ = o.solve(base(), c)
    assert st == "Optimal" and a["Heat"]["calibrated"] and abs(a["Heat"]["drr"] - (heat["drr"] + heat["drr_per_pt_sinter"] * (a["Sinter_share_pct"] - 65.0))) < 1e-9


def test_new_model_constants_reach_the_engine():
    st, _, _, a, _ = o.solve(base(), Config())
    c = cfg(); c.ref_rates["ore"] = 300.0                      # half the reference rate = double the per-kg ore rules
    st2, _, _, a2, _ = o.solve(base(), c)
    assert st2 == "Optimal" and abs(a2["Fuel_terms"]["ore_fe"]) > abs(a["Fuel_terms"]["ore_fe"]) * 1.5
    _, _, _, _, diag = o.solve(base(), cfg(mgo_al2o3_guide_lo=0.60, mgo_al2o3_guide_hi=0.70))
    assert any("MgO/Al2O3" in x and "0.6-0.7" in x for x in diag)
    _, _, _, a3, _ = o.solve(base(), cfg(raw_flux_min_cao_mgo=99.0))
    assert a3["Raw_flux_kg"] == 0.0                            # nothing counts as raw flux any more


def test_heat_balance_is_in_the_bundle_and_the_workbook():
    from openpyxl import load_workbook
    res = o.solve(base(), Config())
    b = o.make_bundle(base(), res, Config())
    T = {k: t for k, _t, t in b["tables"]}
    assert "heat" in T and "heat_rules" in T and "Heat-min coke kg" in b["curve"].columns and b["curve"].attrs.get("heat_note")
    wb = load_workbook(io.BytesIO(o.export_bytes(b)))
    assert "Heat Balance" in wb.sheetnames and len(wb["Heat Balance"]._charts) == 1
    vals = {wb["Run Settings"].cell(r, 2).value for r in range(1, wb["Run Settings"].max_row + 1)}
    assert "Lower-furnace heat losses (MJ/tHM)" in vals and "Mn reduction efficiency" in vals
