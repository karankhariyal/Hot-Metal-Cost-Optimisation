"""
Tests for analytics.py (V6 dashboard analyses).  Run with:  pytest -q tests/test_analytics.py
Every analysis re-solves the unchanged v11.7 engine, so these tests check the arithmetic around the engine and that
no analysis leaves a setting behind in the engine.
"""
import copy
import io
import os
import sys

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from mbf import analytics as an
from mbf import optimiser as o
from mbf.optimiser import Config


@pytest.fixture(scope="module")
def run():
    df, cfg = o.demo_df(), Config()
    res = o.solve(df, cfg)
    return df, cfg, res, o.make_bundle(df, res, cfg)


def engine_defaults_intact():
    return (o.SINTER_MANUAL == {"on": False, "pct": 70.0} and (o.BASICITY_MIN, o.BASICITY_MAX) == (0.99, 1.01)
            and o.AL2O3_MAX_PCT == 18.5 and all(o.FUEL_TERMS.values()) and o.MATERIAL_RULES["ore_fe"]["coef"] == 3.0
            and o.REQUIRED_FE_KGTHM == 965.0)


# ---------------------------------------------------------------- slag oxides
def test_oxide_by_group_adds_up_to_the_engine_totals(run):
    _df, _cfg, res, b = run
    ox = an.oxide_table(b)
    g = an.oxide_by_group(ox)
    tot = g[g["Group"] == "TOTAL CHARGED"].iloc[0]
    body = g[g["Group"] != "TOTAL CHARGED"]
    eng = ox[ox["Material"] == "TOTAL CHARGED"].iloc[0]
    for oxd in an.OXIDE_LIST:
        assert abs(body[f"{oxd} kg"].sum() - eng[f"{oxd} kg"]) < 0.01
        assert abs(body[f"{oxd} % of total"].sum() - 100.0) < 0.1
        assert tot[f"{oxd} % of total"] == 100.0


def test_oxide_cards_match_the_run(run):
    _df, cfg, res, b = run
    cards = {c["Oxide"]: c for c in an.oxide_cards(b)}
    a = res[3]
    assert abs(cards["Al2O3"]["Slag %"] - a["Al2O3_pct"]) < 0.01 and cards["Al2O3"]["High"] == cfg.al2o3_max_pct
    assert abs(cards["SiO2"]["Slag kg"] - a["slag_SiO2_kg"]) < 0.01 and cards["SiO2"]["Charged kg"] > cards["SiO2"]["Slag kg"]
    assert cards["CaO"]["Limit"].startswith("B2") and cards["CaO"]["Status"] in ("OK", "AT LIMIT")


def test_oxide_per_tonne_of_fe(run):
    df, cfg, res, _b = run
    pf = an.oxide_per_fe(df, cfg, res[1]).set_index("Material")
    assert abs(pf.loc["Ore1", "SiO2 kg per t Fe"] - 1000 * 4.0 / 62.0) < 1e-9
    assert abs(pf.loc["Ore1", "Price Rs per t Fe"] - 6300.0 / 0.62) < 1e-6
    assert pf["SiO2 + Al2O3 kg per t Fe"].is_monotonic_increasing          # cleanest first
    assert engine_defaults_intact()


def test_oxides_across_sinter_share(run):
    df, cfg, _res, _b = run
    sw, oxl = an.oxides_across_sinter(df, cfg, 10.0)
    assert list(sw["Sinter share %"].astype(float)) == [float(x) for x in range(0, 101, 10)]
    assert "optimizer's choice" not in set(oxl["Point"])
    assert (sw["Status"] == "Infeasible").iloc[-1]                          # 100 % sinter is blocked on the demo data
    assert engine_defaults_intact()


def test_oxide_whatif_changes_only_the_named_assay(run):
    df, cfg, res, _b = run
    before = df.copy()
    out, status, notes = an.oxide_whatif(df, cfg, res, "Ore1", "Al2O3", 1.0)
    assert status == "Optimal" and out.attrs["change"] == "Ore1 Al2O3: 3.5 % -> 4.5 %"
    assert df.equals(before)                                                # the caller's table is untouched
    row = out.set_index("Result").loc["Cost, Rs/tHM"]
    assert row["What-if"] > row["This run"]                                 # dirtier ore costs more
    with pytest.raises(ValueError):
        an.oxide_whatif(df, cfg, res, "Nope", "Al2O3", 1.0)


# ---------------------------------------------------------------- trends
def test_fe_trend_slope_and_direction(run):
    df, cfg, res, _b = run
    tbl, sl = an.fe_trend(df, cfg, "Iron_ore")
    ok = tbl[tbl["Status"] == "Optimal"]
    assert len(ok) == 7 and ok["Coke kg"].is_monotonic_decreasing and sl["coke"] < 0 and sl["cost"] < 0
    held, sl2 = an.fe_trend(df, cfg, "Iron_ore", hold_share=res[3]["Sinter_share_pct"])
    assert (held.loc[held["Status"] == "Optimal", "Sinter %"] - res[3]["Sinter_share_pct"]).abs().max() < 1e-3
    with pytest.raises(ValueError):
        an.fe_trend(df, cfg, "Minor")                                       # nothing switched on
    assert engine_defaults_intact()


def test_heatmap_cells_equal_pinned_solves(run):
    df, cfg, _res, _b = run
    hm = an.heatmap(df, cfg, "Iron_ore", shares=[60.0, 70.0], deltas=(-1, 0, 1))
    assert len(hm) == 6 and set(hm["Status"]) == {"Optimal"}
    cell = hm[(hm["Sinter %"] == 70.0) & (hm["Fe change (pts)"] == 1.0)].iloc[0]
    d = an._shift(df, "Iron_ore", "Fe", delta=1.0)
    r = o.solve(d, an.pinned(cfg, 70.0), explain=False)
    assert abs(cell["Cost Rs/tHM"] - r[2]) < 1e-6 and abs(cell["Coke kg"] - r[3]["Coke_kg"]) < 1e-6
    assert len(hm.attrs["best"]) == 3
    assert engine_defaults_intact()


def test_price_scan_finds_where_an_ore_enters_the_burden():
    df, cfg = o.demo_df(), Config(stock_balance=0.0)
    df.loc["Ore2"] = df.loc["Ore1"]
    df.loc["Ore2", ["Available", "Price_Rs_t"]] = [True, 6400.0]            # same ore, dearer: not used at today's price
    assert o.solve(df, cfg, explain=False)[1]["Ore2"] < 1e-6
    ps = an.ore_price_scan(df, cfg, "Ore2", pcts=(-10, -5, 0, 5))
    sw = ps.attrs["switches"]
    assert len(sw) == 1 and abs(sw[0]["Price Rs/t"] - 6300.0) < 10.0 and "leaves" in sw[0]["What happens"]
    with pytest.raises(ValueError):
        an.ore_price_scan(o.demo_df(), Config(), "Ore2")                    # switched off


def test_marginal_effects_signs_and_rule_terms(run):
    df, cfg, res, _b = run
    me = an.marginal_effects(df, cfg, res).set_index("Key")
    assert list(me.index) == ["ore_fe", "sinter_fe", "sinter_share", "ore_moisture", "coke_moisture", "coke_price"]
    assert me.loc["ore_fe", "Coke kg"] < 0 and me.loc["sinter_fe", "Coke kg"] < 0
    assert me.loc["ore_moisture", "Coke kg"] > 0 and me.loc["coke_moisture", "Coke kg"] > 0
    a = res[3]
    assert abs(me.loc["ore_fe", "Rule term alone, kg fuel"] - (-3.0 * a["Ore_kg"] / 600.0)) < 1e-9
    assert abs(me.loc["coke_price", "Cost Rs/tHM"] - a["Coke_kg"]) < 0.02 * a["Coke_kg"]      # 1,000 Rs/t x coke tonnes
    assert me.loc["sinter_share", "Status"] == "Optimal"                     # +10 is infeasible here, so it measures -10
    assert engine_defaults_intact()


def test_sinter_share_rule_alone_gives_exactly_ten_kg():
    df = o.demo_df()
    cfg = Config(fuel_terms={k: (k == "sinter_share") for k in o.RULE_SWITCHES})
    res = o.solve(df, cfg)
    me = an.marginal_effects(df, cfg, res).set_index("Key")
    assert abs(me.loc["sinter_share", "Fuel kg"] - (-10.0)) < 1e-6


def test_sinter_basicity_mass_balance():
    df = o.demo_df()
    d, tab = an.sinter_basicity_df(df, 2.3, dilute=True)
    r = tab.iloc[0]
    assert abs(r["B2 what-if"] - 2.3) < 1e-9 and r["Fe % what-if"] < r["Fe % now"]
    old = df.loc["Sinter1", ["Fe", "CaO", "MgO", "SiO2", "Al2O3"]].astype(float)
    new = d.loc["Sinter1", ["Fe", "CaO", "MgO", "SiO2", "Al2O3"]].astype(float)
    added = 2.3 * old["SiO2"] - old["CaO"]
    assert abs(new["Fe"] - old["Fe"] * 100 / (100 + added)) < 1e-9          # diluted in proportion
    d2, _ = an.sinter_basicity_df(df, 2.3, dilute=False)
    assert d2.loc["Sinter1", "Fe"] == df.loc["Sinter1", "Fe"]
    assert o.solve(d, Config())[0] == "Optimal"
    with pytest.raises(ValueError):
        an.sinter_basicity_df(df, 0.0)


# ---------------------------------------------------------------- history, scenarios, export
def test_describe_changes():
    df, cfg = o.demo_df(), Config()
    assert an.describe_changes(df, df.copy(), cfg, copy.deepcopy(cfg)) == []
    d2, c2 = df.copy(), Config(sinter_max=0.72)
    d2.loc["Sinter1", "Price_Rs_t"] = 5200.0
    d2.loc["Limestone", "Available"] = False
    c2.fuel_terms["slag"] = False
    ch = an.describe_changes(df, d2, cfg, c2)
    assert "Sinter top % 80 -> 72" in ch and "Sinter1 price 5,500 -> 5,200" in ch
    assert "Limestone switched off" in ch and "Slag volume rule off" in ch
    assert an.change_label(ch, 2).endswith("(+2 more)")
    c3 = Config(); c3.heat["mode"] = "floor"; c3.heat["loss_MJ_tHM"] = 650.0; c3.ref_rates["coke"] = 410.0
    c3.furnace_profiles["MBF-2"]["base"] = 540.0; c3.mn_reduction_eff = 0.8
    ch3 = an.describe_changes(df, df.copy(), cfg, c3)
    assert "Heat balance mode check -> floor" in ch3 and any(x.startswith("Lower-furnace heat losses") for x in ch3)
    assert "reference coke rate 402 -> 410 kg/tHM" in ch3 and "MBF-2 base fuel 545 -> 540" in ch3
    assert "Mn reduction efficiency 0.85 -> 0.8" in ch3


def test_scenarios_compare(run):
    df, cfg, res, b = run
    s1 = an.scenario_from_bundle("Today", b)
    c2 = Config(sinter_manual_on=True, sinter_manual_pct=65.0)
    r2 = o.solve(df, c2)
    s2 = an.scenario_from_bundle("Pinned 65", o.make_bundle(df, r2, c2, with_curve=False))
    kp, rc, diffs = an.compare_scenarios([s1, s2])
    kp = kp.set_index("Result")
    assert abs(kp.loc["Cost Rs/tHM", "Pinned 65 vs Today"] - (r2[2] - res[2])) < 1e-9
    assert set(rc["Material"]) >= {"Ore1", "Sinter1", "Coke3"}
    assert diffs[0][2] == ["Sinter pin off -> on", "Pinned sinter % 70 -> 65"]
    with pytest.raises(ValueError):
        an.scenario_from_bundle("bad", o.make_bundle(df, o.solve(df, Config(sinter_manual_on=True, sinter_manual_pct=99.0)), Config()))


def test_export_with_dashboard_sheets_keeps_the_engine_workbook(run):
    from openpyxl import load_workbook
    df, cfg, res, b = run
    me = an.marginal_effects(df, cfg, res)
    extra = an.extra_sheets([("Rules of Thumb", "Rules of thumb", me, "note"), ("Empty", "nothing", pd.DataFrame(), "")])
    wb = load_workbook(io.BytesIO(o.export_bytes(b, {}, extra=extra)))
    assert "Rules of Thumb" in wb.sheetnames and "Empty" not in wb.sheetnames
    assert wb.sheetnames[-1] == "Run Settings"                              # engine order kept
    assert len(wb["Fuel Rate"]._charts) >= 1 and len(wb["Sinter Curve 0-100%"]._charts) == 2
    ws = wb["Rules of Thumb"]
    assert ws.cell(4, 1).value == "Key" and ws.max_row == 4 + len(me)
    plain = load_workbook(io.BytesIO(o.export_bytes(b, {})))
    assert "Rules of Thumb" not in plain.sheetnames                          # no hook = the V5 workbook
