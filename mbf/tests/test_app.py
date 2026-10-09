"""
Dashboard tests: drive app.py the way a user would (Streamlit's AppTest).  Run with:  pytest -q tests/test_app.py
File uploads and the data editors cannot be driven by AppTest, so those paths are covered by setting session state
directly and by the engine tests (loader, template round trip).
"""
import dataclasses
import io
import os
import sys

import pytest
from streamlit.testing.v1 import AppTest

ROOT = os.path.join(os.path.dirname(__file__), "..")
sys.path.insert(0, ROOT)
from mbf import optimiser as opt  # noqa: E402

PAGES = ["Dashboard", "Inputs", "Materials & stock", "Burden & cost", "Slag oxides", "Trends", "Scenario analysis",
         "Reports & export", "Heat audit", "Upload & settings"]


def new_app():
    at = AppTest.from_file(os.path.join(ROOT, "app.py"), default_timeout=120)
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def goto(at, page):
    at.session_state["nav"] = page
    at.run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def run_opt(at):
    goto(at, "Dashboard")
    at.button(key="run_opt").click().run()
    assert not at.exception, [e.value for e in at.exception]
    return at


def text_of(at):
    return " ".join(m.value for m in at.markdown)


def test_every_page_loads_before_and_after_a_run():
    at = new_app()
    for p in PAGES:                                  # before any run: pages that need a result must say so, not crash
        goto(at, p)
    run_opt(at)
    assert at.session_state["result"][0] == "Optimal"
    for p in PAGES:
        goto(at, p)


def test_dashboard_starts_with_demo_notice_and_a_run_button():
    at = new_app()
    assert "Demo table in use" in text_of(at)
    assert not at.button(key="run_opt").disabled


def test_run_shows_the_new_v115_sections():
    at = run_opt(new_app())
    txt = text_of(at)
    for k in ("What set the sinter share", "Coke and cost at every sinter share", "Where the CaO, SiO2 and Al2O3 come from"):
        assert k in txt, k
    b = at.session_state["bundle"]
    assert b["curve"] is not None and len(b["curve"]) == 41 and b["decision"].startswith("SINTER DECISION")


def test_run_shows_the_engine_answer():
    at = run_opt(new_app())
    cost = at.session_state["result"][2]
    assert f"{cost:,.0f}" in text_of(at)
    assert abs(cost - opt.solve(opt.demo_df(), opt.Config())[2]) < 1e-9
    assert "Up to date, run 1" in text_of(at)
    assert at.session_state["bundle"]["status"] == "Optimal"


def test_changing_an_input_asks_for_a_rerun_and_changes_the_result():
    at = run_opt(new_app())
    goto(at, "Inputs")
    at.number_input(key="cfg_pci_fixed_kgthm_0").set_value(100.0).run()
    assert at.session_state["cfg"].pci_fixed_kgthm == 100.0
    assert at.session_state["changed"] is True
    goto(at, "Dashboard")
    assert "Rerun needed: Inputs changed" in text_of(at)
    at.button(key="run_opt").click().run()
    a = at.session_state["result"][3]
    assert abs(a["PCI_kg"] - 100.0) < 1e-6
    assert at.session_state["runs"] == 2 and len(at.session_state["history"]) == 2


def test_percent_inputs_are_converted_to_fractions():
    at = new_app()
    goto(at, "Inputs")
    at.number_input(key="cfg_sinter_max_0").set_value(72.0).run()
    assert abs(at.session_state["cfg"].sinter_max - 0.72) < 1e-12
    run_opt(at)
    assert at.session_state["result"][3]["Sinter_share_pct"] <= 72.0 + 1e-3


def test_pin_the_sinter_share():
    at = new_app()
    goto(at, "Inputs")
    at.checkbox(key="cfg_sinter_manual_on_0").check().run()
    at.number_input(key="cfg_sinter_manual_pct_0").set_value(68.0).run()
    run_opt(at)
    assert abs(at.session_state["result"][3]["Sinter_share_pct"] - 68.0) < 1e-3
    assert "Pinned at 68 %" in text_of(at)


def test_conflicting_settings_block_the_run():
    at = new_app()
    goto(at, "Inputs")
    at.number_input(key="cfg_basicity_min_0").set_value(1.4).run()      # min above max (1.01)
    assert "minimum is above the maximum" in text_of(at)
    goto(at, "Dashboard")
    assert at.button(key="run_opt").disabled
    assert "Fix these before running" in text_of(at)


def test_reset_to_plant_defaults():
    at = new_app()
    goto(at, "Inputs")
    at.number_input(key="cfg_pci_fixed_kgthm_0").set_value(90.0).run()
    at.button(key="reset_cfg").click().run()
    assert at.session_state["cfg"].pci_fixed_kgthm == 120.0
    assert at.number_input(key="cfg_pci_fixed_kgthm_1").value == 120.0


def test_thumb_rule_switches_turn_single_rules_off_and_on():
    at = new_app()
    goto(at, "Inputs")
    assert all(at.session_state["cfg"].fuel_terms.values())                 # every rule on by default
    at.toggle(key="cfg_rule_other_moisture_0").set_value(False).run()
    ft = at.session_state["cfg"].fuel_terms
    assert ft["other_moisture"] is False and ft["ore_moisture"] is True and at.session_state["changed"] is True
    run_opt(at)
    assert any("SWITCHED OFF" in x for x in at.session_state["result"][4])
    assert "Thumb rules switched off" in text_of(at)
    goto(at, "Inputs")
    at.button(key="rules_all_on").click().run()
    assert all(at.session_state["cfg"].fuel_terms.values())


def test_a_failed_run_is_explained_and_other_pages_do_not_crash():
    at = new_app()
    d = opt.demo_df()
    d.loc["Coke3", "Available"] = False
    at.session_state["df"] = d
    run_opt(at)
    assert at.session_state["result"][0] == "NO_PRODUCTION"
    assert "No burden was produced" in text_of(at) and "No regular coke" in text_of(at)
    goto(at, "Burden & cost")
    assert "ended as NO_PRODUCTION" in text_of(at)
    goto(at, "Reports & export")
    assert "inputs and the reason" in text_of(at)


def test_material_table_errors_block_the_run():
    at = new_app()
    d = opt.demo_df()
    d.loc["Sponge_Iron", "Available"] = True                               # on, but no assay and no price
    at.session_state["df"] = d
    goto(at, "Dashboard")
    assert at.button(key="run_opt").disabled and "no price" in text_of(at)


def test_sinter_sweep_and_sensitivity():
    at = run_opt(new_app())
    goto(at, "Scenario analysis")
    at.number_input(key="sw_step").set_value(5.0).run()
    at.button(key="run_sweep").click().run()
    assert not at.exception, [e.value for e in at.exception]
    sw = at.session_state["analysis"]["sweep"]["df"]
    assert len(sw) == 8 and (sw["Status"] == "Optimal").all()                 # free choice + 50, 55 ... 80
    assert len(at.session_state["analysis"]["sweep_oxides"]["df"]) > 0
    at.button(key="run_sens").click().run()
    assert not at.exception, [e.value for e in at.exception]
    sn = at.session_state["analysis"]["sensitivity"]["df"]
    assert len(sn) == 5 and (sn["Status"] == "Optimal").all()


def test_sensitivity_on_a_group_with_nothing_on_is_reported_not_raised():
    at = new_app()
    goto(at, "Scenario analysis")
    at.selectbox(key="sens_group").select("Minor").run()                   # no minor material is on in the demo table
    at.button(key="run_sens").click().run()
    assert not at.exception and "could not run" in text_of(at)


def test_price_tornado_and_breakeven_tools():
    at = run_opt(new_app())
    goto(at, "Scenario analysis")
    at.button(key="run_price").click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert (at.session_state["analysis"]["price"]["df"]["Status"] == "Optimal").all()
    at.button(key="run_tornado").click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.session_state["analysis"]["tornado"]["df"]["Swing Rs"].is_monotonic_decreasing
    at.button(key="run_be").click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert "breakeven" in at.session_state["analysis"]


def test_breakeven_is_blocked_while_the_share_is_pinned():
    at = new_app()
    goto(at, "Inputs")
    at.checkbox(key="cfg_sinter_manual_on_0").check().run()
    goto(at, "Scenario analysis")
    assert "The sinter share is pinned" in text_of(at)


def test_export_is_built_only_on_click_and_is_a_real_workbook():
    from openpyxl import load_workbook
    at = run_opt(new_app())
    goto(at, "Reports & export")
    assert at.session_state["export"] is None                              # nothing is prepared by itself
    assert not [b for b in at.get("download_button")]
    at.button(key="export_build").click().run()
    assert not at.exception, [e.value for e in at.exception]
    data, name, _ = at.session_state["export"]
    assert name.startswith("MBF_Optimised_Results_MBF-3_") and name.endswith(".xlsx")
    wb = load_workbook(io.BytesIO(data))
    assert {"Summary", "Sinter Curve 0-100%", "Inputs", "Optimised Burden", "Run Settings"} <= set(wb.sheetnames)
    assert at.get("download_button")                                       # the download control now exists


def test_export_includes_the_sweep_when_one_was_run():
    from openpyxl import load_workbook
    at = run_opt(new_app())
    goto(at, "Scenario analysis")
    at.number_input(key="sw_step").set_value(5.0).run()
    at.button(key="run_sweep").click().run()
    goto(at, "Reports & export")
    at.button(key="export_build").click().run()
    wb = load_workbook(io.BytesIO(at.session_state["export"][0]))
    assert "Sinter Sweep" in wb.sheetnames and "Oxides by Sinter %" in wb.sheetnames


def test_a_new_run_clears_the_prepared_export():
    at = run_opt(new_app())
    goto(at, "Reports & export")
    at.button(key="export_build").click().run()
    assert at.session_state["export"] is not None
    run_opt(at)
    assert at.session_state["export"] is None


def test_settings_page_self_test_and_demo_switch():
    at = new_app()
    goto(at, "Upload & settings")
    at.button(key="selftest_btn").click().run()
    assert all(ok for _, ok, _ in at.session_state["selftest"])
    at.button(key="hide_demo").click().run()
    assert at.session_state["demo"] is False


def test_two_sessions_do_not_share_settings():
    a = new_app()
    b = new_app()
    goto(a, "Inputs")
    a.number_input(key="cfg_pci_fixed_kgthm_0").set_value(90.0).run()
    run_opt(a)
    run_opt(b)
    assert abs(a.session_state["result"][3]["PCI_kg"] - 90.0) < 1e-6
    assert abs(b.session_state["result"][3]["PCI_kg"] - 120.0) < 1e-6


# ---------------------------------------------------------------- V6
def test_recipe_sits_directly_under_the_key_numbers():
    at = run_opt(new_app())
    txt = text_of(at)
    assert txt.index("Raw-material cost") < txt.index("Recipe") < txt.index("What set the sinter share") < txt.index("Limits")
    assert at.text_input(key="dash_scen_name_1").value == "Run 1"


def test_quick_inputs_panel_applies_runs_and_says_what_changed():
    at = run_opt(new_app())
    at.button(key="qi_open").click().run()
    v = at.session_state["qi_ver"]
    at.number_input(key=f"qi{v}_price_Sinter1").set_value(5000.0).run()
    at.slider(key=f"qi{v}_cfg_al2o3_min_pct").set_value((17.0, 19.0)).run()
    assert at.session_state["df"].loc["Sinter1", "Price_Rs_t"] == 5500.0          # still a draft
    assert at.session_state["cfg"].al2o3_max_pct == 18.5
    assert "Changes to apply" in text_of(at)
    at.button(key=f"qi{v}_apply").click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.session_state["runs"] == 2 and at.session_state["qi_show"] is False
    assert at.session_state["cfg"].al2o3_max_pct == 19.0 and at.session_state["df"].loc["Sinter1", "Price_Rs_t"] == 5000.0
    assert at.session_state["history"][-1]["Changed"] == "Al2O3 max % 18.5 -> 19; Sinter1 price 5,500 -> 5,000"
    assert "vs run 1: cost" in text_of(at)
    goto(at, "Inputs")                                                          # the Inputs page shows the same value
    assert at.number_input(key=f"cfg_al2o3_max_pct_{at.session_state['cfg_ver']}").value == 19.0


def test_quick_inputs_cancel_and_defaults_change_nothing():
    at = run_opt(new_app())
    at.button(key="qi_open").click().run()
    v = at.session_state["qi_ver"]
    at.number_input(key=f"qi{v}_cfg_pci_fixed_kgthm").set_value(90.0).run()
    at.button(key=f"qi{v}_defaults").click().run()
    assert at.session_state["qi_draft"]["cfg"].pci_fixed_kgthm == 120.0
    v = at.session_state["qi_ver"]
    at.button(key=f"qi{v}_cancel").click().run()
    assert at.session_state["qi_draft"] is None and at.session_state["cfg"].pci_fixed_kgthm == 120.0 and at.session_state["runs"] == 1


def test_quick_inputs_refuse_a_broken_table():
    at = run_opt(new_app())
    at.button(key="qi_open").click().run()
    v = at.session_state["qi_ver"]
    at.number_input(key=f"qi{v}_price_Coke3").set_value(0.0).run()
    at.button(key=f"qi{v}_apply").click().run()
    assert "Not applied" in text_of(at) and at.session_state["runs"] == 1


def test_slag_oxides_page_views_and_whatif():
    at = run_opt(new_app())
    goto(at, "Slag oxides")
    txt = text_of(at)
    for k in ("CaO charged", "Al2O3 charged", "Oxide per tonne of Fe delivered", "Across sinter share"):
        assert k in txt, k
    at.radio(key="ox_sel").set_value("Al2O3").run()
    at.radio(key="ox_view").set_value("By group").run()
    assert not at.exception, [e.value for e in at.exception]
    at.selectbox(key="wi_mat").select("Ore1").run()
    at.selectbox(key="wi_ox").select("SiO2").run()
    at.button(key="wi_run").click().run()
    assert not at.exception, [e.value for e in at.exception]
    assert at.session_state["ox_whatif"]["status"] == "Optimal"
    assert "perfe" in at.session_state["v6"]


def test_trends_page_builds_every_analysis():
    at = run_opt(new_app())
    goto(at, "Trends")
    assert "rules of thumb on this run" in text_of(at)
    v6 = at.session_state["v6"]
    for k in ("marginal", "fe_trend", "heatmap", "price_scan"):
        assert k in v6, k
    assert "basicity" not in v6                              # nothing to test at today's sinter basicity
    at.checkbox(key="tr_fe_hold").check().run()
    at.radio(key="hm_span").set_value("0 to 100 %").run()
    at.slider(key="sb_target").set_value(2.3).run()
    assert not at.exception, [e.value for e in at.exception]
    assert "B2 2.30" in at.session_state["v6"]["basicity"]["title"]
    assert "held" in at.session_state["v6"]["fe_trend"]["title"]


def test_scenarios_save_overlay_compare_and_export():
    from openpyxl import load_workbook
    at = run_opt(new_app())
    at.button(key="dash_scen_save").click().run()
    goto(at, "Inputs")
    at.number_input(key="cfg_sinter_max_0").set_value(70.0).run()
    run_opt(at)
    at.text_input(key="dash_scen_name_2").set_value("Top 70").run()
    at.button(key="dash_scen_save").click().run()
    assert [s["name"] for s in at.session_state["scenarios"]] == ["Run 1", "Top 70"]
    goto(at, "Trends")
    at.multiselect(key="tr_overlay").set_value(["Run 1", "Top 70"]).run()
    assert not at.exception, [e.value for e in at.exception]
    assert "Top 70 vs Run 1: Sinter top % 80 -> 70" in text_of(at)
    goto(at, "Slag oxides")
    goto(at, "Reports & export")
    at.button(key="export_build").click().run()
    wb = load_workbook(io.BytesIO(at.session_state["export"][0]))
    for sh in ("Rules of Thumb", "Oxides per t Fe", "Fe Trend", "Heat Map", "Run History", "Scenario Comparison", "Scenario Recipes"):
        assert sh in wb.sheetnames, sh
    assert wb.sheetnames[-1] == "Run Settings"


def test_analysis_pages_warn_when_inputs_changed_after_the_run():
    at = run_opt(new_app())
    goto(at, "Inputs")
    at.number_input(key="cfg_pci_fixed_kgthm_0").set_value(110.0).run()
    goto(at, "Trends")
    assert "These analyses use the last run" in text_of(at)


# ---------------------------------------------------------------- V7: every input in the panel, heat balance, heat audit
def _panel_keys(at):
    v = at.session_state["qi_ver"]
    p = f"qi{v}_"
    keys = set()
    for kind in ("number_input", "slider", "selectbox", "toggle", "checkbox", "text_input"):
        for w in getattr(at, kind):
            if w.key and w.key.startswith(p):
                keys.add(w.key[len(p):])
    return keys


def test_quick_inputs_panel_covers_every_model_input():
    at = run_opt(new_app())
    at.button(key="qi_open").click().run()
    keys = _panel_keys(at)
    c = opt.Config()
    two_ended = {"sinter_max": "cfg_sinter_min", "basicity_max": "cfg_basicity_min", "mgo_max_pct": "cfg_mgo_min_pct",
                 "al2o3_max_pct": "cfg_al2o3_min_pct"}                       # range sliders: one widget sets both ends
    missing = []
    for f in dataclasses.fields(c):
        n = f.name
        need = {"fuel_terms": {f"rule_{k}" for k in opt.RULE_SWITCHES},
                "material_rules": {f"rule_{k}_{p}" for k in opt.MATERIAL_RULES for p in ("coef", "ref")},
                "furnace_profiles": {f"base_{fn}" for fn in c.furnace_profiles},
                "ref_rates": {f"ref_{k}" for k in c.ref_rates},
                "heat": {"heat_mode"} | {f"heat_{k}" for k in opt.HEAT_FIELDS}}.get(n, {two_ended.get(n, f"cfg_{n}")})
        missing += sorted(need - keys)
    df = at.session_state["df"]
    for m in df.index:                                                      # on/off and price of EVERY material, on or off
        missing += sorted({f"on_{m}", f"price_{m}"} - keys)
    m = at.session_state["qi_draft"]["df"].index[0]                         # the detail editor: every other column
    cols = set(opt.COLUMNS[1:]) - {"Available", "Price_Rs_t"}
    missing += sorted(({f"mat_{m}_{c_}" for c_ in cols - set(opt.OPT_COLS)} | {f"mat_{m}_{c_}_on" for c_ in opt.OPT_COLS}) - keys)
    assert not missing, missing


def test_quick_inputs_edit_materials_heat_and_constants():
    at = run_opt(new_app())
    at.button(key="qi_open").click().run()
    v = at.session_state["qi_ver"]
    at.selectbox(key=f"qi{v}_mat_pick").set_value("Coke3").run()
    at.number_input(key=f"qi{v}_mat_Coke3_FC").set_value(87.0).run()
    at.toggle(key=f"qi{v}_on_Dolomite").set_value(False).run()
    at.number_input(key=f"qi{v}_heat_blast_temp_C").set_value(1100.0).run()
    at.number_input(key=f"qi{v}_ref_coke").set_value(410.0).run()
    at.number_input(key=f"qi{v}_cfg_mn_reduction_eff").set_value(0.8).run()
    assert at.session_state["cfg"].heat["blast_temp_C"] == 1000.0                   # still a draft
    at.button(key=f"qi{v}_apply").click().run()
    assert not at.exception, [e.value for e in at.exception]
    c, d = at.session_state["cfg"], at.session_state["df"]
    assert c.heat["blast_temp_C"] == 1100.0 and not c.heat["calibrated"] and c.ref_rates["coke"] == 410.0 and c.mn_reduction_eff == 0.8
    assert d.loc["Coke3", "FC"] == 87.0 and not bool(d.loc["Dolomite", "Available"]) and at.session_state["runs"] == 2


def test_quick_inputs_add_a_material():
    at = run_opt(new_app())
    at.button(key="qi_open").click().run()
    v = at.session_state["qi_ver"]
    at.text_input(key=f"qi{v}_new_name").set_value("Ore 3").run()
    at.button(key=f"qi{v}_new_add").click().run()
    d = at.session_state["qi_draft"]["df"]
    assert "Ore_3" in d.index and not bool(d.loc["Ore_3", "Available"])
    assert "Ore_3 added" in text_of(at)


def test_calibrate_losses_from_the_panel_closes_the_balance():
    at = run_opt(new_app())
    at.button(key="qi_open").click().run()
    v = at.session_state["qi_ver"]
    at.button(key=f"qi{v}_heat_calib").click().run()
    assert at.session_state["qi_draft"]["cfg"].heat["calibrated"]
    v = at.session_state["qi_ver"]
    at.button(key=f"qi{v}_apply").click().run()
    assert not at.exception, [e.value for e in at.exception]
    a = at.session_state["result"][3]
    assert abs(a["Heat"]["C_surplus_kg"]) < 1e-6 and at.session_state["cfg"].heat["calibrated"]


def test_inputs_heat_tab_floor_mode_and_heat_views():
    at = new_app()
    goto(at, "Inputs")
    at.selectbox(key="cfg_heat_mode_0").set_value("floor").run()
    assert at.session_state["cfg"].heat["mode"] == "floor" and at.session_state["changed"]
    run_opt(at)
    diag = at.session_state["result"][4]
    assert any(x.startswith("HEAT CHECK") for x in diag)
    assert "Heat balance, hot zone" in text_of(at)
    goto(at, "Burden & cost")
    assert "Hot-zone heat demand" in text_of(at) and "Heat balance against the plant thumb rules" in text_of(at)


def test_heat_audit_page_runs_and_applies():
    at = new_app()
    goto(at, "Heat audit")
    at.button(key="aud_demo").click().run()
    assert len(at.session_state["audit_recs"]) == 4
    at.button(key="aud_run").click().run()
    assert not at.exception, [e.value for e in at.exception]
    A = at.session_state["audit"]
    assert A is not None and A["fits"]["drr"]["applied"]
    at.button(key="aud_apply").click().run()
    assert not at.exception, [e.value for e in at.exception]
    h = at.session_state["cfg"].heat
    assert h["calibrated"] and abs(h["drr_per_pt_sinter"] + 0.004) < 4e-4 and at.session_state["changed"]
    assert "Applied to the model" in text_of(at)
