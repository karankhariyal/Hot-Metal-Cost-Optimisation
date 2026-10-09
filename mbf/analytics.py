"""
MBF burden optimiser - dashboard ANALYTICS (V6, carried into V7)
================================================================
Extra calculations for the V6 pages (Slag oxides, Trends, quick inputs, saved scenarios, run history).

Nothing here changes the model. Every function re-solves the unchanged v11.7 engine in optimiser.py, and every
solve goes through `optimiser.session(cfg)` (directly or through `optimiser.solve`), so one user's settings can never
leak into another user's run and the engine globals are always restored afterwards.
"""
import copy
import dataclasses
import math

import numpy as np
import pandas as pd

from mbf import optimiser as opt

OXIDE_LIST = ["CaO", "SiO2", "Al2O3", "MgO"]
_NOT_BODY = ("TOTAL CHARGED", "Less: SiO2 reduced to Si in hot metal", "SLAG")


# ================================================================ keys and small helpers
def cfg_key(cfg):
    """A stable text key for a Config (used to cache results per set of settings)."""
    return repr(sorted(dataclasses.asdict(cfg).items()))


def run_key(df, cfg, *extra):
    return (opt.fingerprint(opt.ensure_columns(df)), cfg_key(cfg)) + tuple(extra)


def pinned(cfg, pct):
    """A copy of cfg with the sinter share pinned at pct (%)."""
    c = copy.deepcopy(cfg)
    c.sinter_manual_on = True
    c.sinter_manual_pct = float(pct)
    return c


def on_materials(df, group):
    d = opt.ensure_columns(df)
    return [m for m in d.index if str(d.loc[m, "Group"]).strip() == group and bool(d.loc[m, "Available"])]


def summary(res):
    """Headline numbers of one solve result (status, blend, cost, achieved, diagnostics)."""
    status, blend, cost, ach, _diag = res
    if status != "Optimal":
        return {"Status": status}
    return {"Status": status, "Cost Rs/tHM": cost, "Sinter %": ach["Sinter_share_pct"], "Coke kg": ach["Coke_kg"],
            "Fuel kg": ach["Fuel_supplied"], "Slag kg": ach["slag_kg"], "B2": ach["B2"], "Al2O3 %": ach["Al2O3_pct"],
            "MgO %": ach["MgO_pct"], "Raw flux kg": ach["Raw_flux_kg"], "Flux kg": ach["Flux_kg"],
            "Ore kg": ach["Ore_kg"], "Sinter kg": ach["Sinter_kg"]}


def _shift(df, group, column, delta=None, factor=None, add=None):
    """Copy of df with `column` of every ON material in `group` changed (delta points, x factor, or + add)."""
    d = opt._fdf(opt.ensure_columns(df))
    for m in on_materials(d, group):
        v = float(d.loc[m, column])
        if delta is not None:
            v = v + float(delta)
        if factor is not None:
            v = v * float(factor)
        if add is not None:
            v = v + float(add)
        d.loc[m, column] = max(0.0, v)
    return d


def _slope(x, y):
    x, y = np.asarray(x, float), np.asarray(y, float)
    ok = np.isfinite(x) & np.isfinite(y)
    if ok.sum() < 2 or np.ptp(x[ok]) == 0:
        return None
    return float(np.polyfit(x[ok], y[ok], 1)[0])


# ================================================================ SLAG OXIDES page
def oxide_table(bundle):
    """The engine's SLAG OXIDE SOURCES table of this run."""
    return {k: t for k, _t, t in bundle["tables"]}["oxides"]


def oxide_body(ox):
    return ox[~ox["Material"].isin(_NOT_BODY)].reset_index(drop=True)


def oxide_by_group(ox):
    """kg and % of each oxide charged, summed by material group (groups in the engine's order)."""
    body = oxide_body(ox)
    cols = [f"{o} kg" for o in OXIDE_LIST] + ["kg/tHM charged"]
    g = body.groupby("Group", sort=False)[cols].sum()
    order = {x: i for i, x in enumerate(opt.GROUPS)}
    g = g.loc[sorted(g.index, key=lambda x: order.get(str(x), 99))]
    tot = {o: float(body[f"{o} kg"].sum()) for o in OXIDE_LIST}
    out = g.reset_index()
    for o in OXIDE_LIST:
        out.insert(out.columns.get_loc(f"{o} kg") + 1, f"{o} % of total",
                   100.0 * out[f"{o} kg"] / tot[o] if tot[o] > 1e-12 else 0.0)
    total = {"Group": "TOTAL CHARGED", "kg/tHM charged": float(body["kg/tHM charged"].sum())}
    for o in OXIDE_LIST:
        total[f"{o} kg"] = tot[o]
        total[f"{o} % of total"] = 100.0
    return pd.concat([out, pd.DataFrame([total])], ignore_index=True)


def oxide_cards(bundle):
    """One card per oxide: kg charged, kg in slag, % of slag, the limit that governs it and where it sits."""
    ox, ach, c = oxide_table(bundle), bundle["ach"], bundle["cfg"]
    tot = ox[ox["Material"] == "TOTAL CHARGED"].iloc[0]
    slag = ox[ox["Material"] == "SLAG"].iloc[0]
    b2 = ("B2 = CaO / SiO2", ach["B2"], c.basicity_min, c.basicity_max, "", 3)
    limits = {"CaO": b2, "SiO2": b2,
              "Al2O3": ("Al2O3 in slag", ach["Al2O3_pct"], c.al2o3_min_pct, c.al2o3_max_pct, " %", 2),
              "MgO": ("MgO in slag", ach["MgO_pct"], c.mgo_min_pct, c.mgo_max_pct, " %", 2)}
    cards = []
    for o in OXIDE_LIST:
        lab, v, lo, hi, unit, dec = limits[o]
        cards.append({"Oxide": o, "Charged kg": float(tot[f"{o} kg"]), "Slag kg": float(slag[f"{o} kg"]),
                      "Slag %": float(slag[f"{o} % of total"]), "Limit": lab, "Value": float(v), "Low": float(lo),
                      "High": float(hi), "Unit": unit, "Dec": dec, "Status": opt.band_status(float(v), float(lo), float(hi))})
    return cards


def oxide_per_fe(df, cfg, blend=None):
    """For every iron-bearing material switched on: kg of each oxide it brings per tonne of Fe it delivers,
    and its price per tonne of contained Fe. Sorted cleanest first (least SiO2 + Al2O3 per tonne of Fe)."""
    d = opt.ensure_columns(df)
    rows = []
    with opt.session(cfg):
        for g in ("Iron_ore", "Sinter", "Minor"):
            for m in on_materials(d, g):
                fe = float(d.loc[m, "Fe"])
                if fe <= 0:
                    continue
                r = {"Material": m, "Group": g, "Used kg/tHM": float((blend or {}).get(m, 0.0)) if blend is not None else np.nan,
                     "Fe %": fe}
                for o in OXIDE_LIST:
                    r[f"{o} kg per t Fe"] = 1000.0 * float(d.loc[m, o]) / fe
                r["SiO2 + Al2O3 kg per t Fe"] = r["SiO2 kg per t Fe"] + r["Al2O3 kg per t Fe"]
                r["Price Rs per t Fe"] = opt.eff_price(d, m) / (fe / 100.0)
                rows.append(r)
    out = pd.DataFrame(rows)
    if len(out):
        out = out.sort_values("SiO2 + Al2O3 kg per t Fe").reset_index(drop=True)
    return out


def oxides_across_sinter(df, cfg, step=5.0):
    """Model re-solved with the sinter share pinned from 0 % to 100 %: (sweep table, long table of oxide kg by material).
    The optimiser's-choice row is removed from both."""
    out, oxl = opt.sweep(df, cfg, 0.0, 100.0, float(step), with_oxides=True)
    out = out[out["Sinter share %"] != "optimizer's choice"].reset_index(drop=True)
    if len(oxl):
        oxl = oxl[oxl["Point"] != "optimizer's choice"].reset_index(drop=True)
    return out, oxl


def oxide_whatif(df, cfg, base_res, material, oxide, delta):
    """Change one oxide of one material by `delta` points and re-solve under the same settings.
    Returns (comparison table, status of the what-if, notes)."""
    d = opt._fdf(opt.ensure_columns(df))
    if material not in d.index:
        raise ValueError(f"Unknown material '{material}'.")
    if oxide not in OXIDE_LIST:
        raise ValueError(f"Oxide must be one of {OXIDE_LIST}.")
    old = float(d.loc[material, oxide])
    d.loc[material, oxide] = max(0.0, old + float(delta))
    res = opt.solve(d, cfg, explain=False)
    notes = []
    if res[0] != "Optimal":
        res = opt.solve(d, cfg, explain=True)
        notes = list(res[4])
    a0, a1 = base_res[3], (res[3] if res[0] == "Optimal" else None)
    metrics = [("Slag Al2O3, %", "Al2O3_pct", 2), ("Slag MgO, %", "MgO_pct", 2), ("Basicity B2", "B2", 3),
               ("Slag, kg/tHM", "slag_kg", 1), ("Regular coke, kg/tHM", "Coke_kg", 1), ("Total flux, kg/tHM", "Flux_kg", 1),
               ("Sinter share, %", "Sinter_share_pct", 2), ("Cost, Rs/tHM", "Cost_Rs_tHM", 0)]
    rows = []
    for lab, k, dec in metrics:
        v0 = float(a0[k])
        v1 = float(a1[k]) if a1 is not None else np.nan
        rows.append({"Result": lab, "This run": round(v0, dec), "What-if": round(v1, dec) if a1 is not None else np.nan,
                     "Change": round(v1 - v0, dec) if a1 is not None else np.nan})
    out = pd.DataFrame(rows)
    out.attrs["change"] = f"{material} {oxide}: {old:g} % -> {float(d.loc[material, oxide]):g} %"
    return out, res[0], notes


# ================================================================ TRENDS page
def fe_trend(df, cfg, group="Iron_ore", deltas=(-3, -2, -1, 0, 1, 2, 3), hold_share=None):
    """Fe of every ON material in `group` shifted by each delta and the model re-solved.
    hold_share: pin the sinter share at this % so only the Fe effect shows; None = the settings as applied.
    Returns (table, slopes {'coke': kg per pt, 'cost': Rs per pt, 'fuel': kg per pt})."""
    if not on_materials(df, group):
        raise ValueError(f"No {group} material is switched on.")
    c = pinned(cfg, hold_share) if hold_share is not None else cfg
    with opt.session(c):
        out = opt.assay_sensitivity(opt.ensure_columns(df), group, "Fe", tuple(float(x) for x in deltas))
    ok = out[out["Status"] == "Optimal"]
    slopes = {"coke": _slope(ok["Change (pts)"], ok["Coke kg"]) if len(ok) else None,
              "fuel": _slope(ok["Change (pts)"], ok["Fuel kg"]) if len(ok) else None,
              "cost": _slope(ok["Change (pts)"], ok["Cost Rs/tHM"]) if len(ok) else None}
    return out, slopes


def heatmap(df, cfg, group="Iron_ore", shares=None, deltas=(-3, -2, -1, 0, 1, 2, 3)):
    """Cost and coke over a grid of pinned sinter share x Fe shift of `group`. Long table; infeasible cells are kept
    with blank results. attrs['best']: least-cost sinter share for each Fe shift."""
    if not on_materials(df, group):
        raise ValueError(f"No {group} material is switched on.")
    if shares is None:
        shares = np.arange(cfg.sinter_min * 100, cfg.sinter_max * 100 + 1e-9, 5.0)
    shares = [float(s) for s in shares]
    deltas = [float(x) for x in deltas]
    if len(shares) * len(deltas) > 400:
        raise ValueError("Grid limited to 400 cells - use a larger step.")
    rows = []
    with opt.session(cfg):
        for dl in deltas:
            d = _shift(df, group, "Fe", delta=dl)
            for s in shares:
                opt.SINTER_MANUAL.update(on=True, pct=s)
                st, _b, c, a, _ = opt.solve_mbf(d, explain=False)
                r = {"Fe change (pts)": dl, "Sinter %": s, "Status": st}
                if st == "Optimal":
                    r.update({"Cost Rs/tHM": c, "Coke kg": a["Coke_kg"], "Fuel kg": a["Fuel_supplied"]})
                rows.append(r)
    out = pd.DataFrame(rows)
    for c_ in ("Cost Rs/tHM", "Coke kg", "Fuel kg"):
        if c_ not in out.columns:
            out[c_] = np.nan
    best = []
    for dl in deltas:
        sub = out[(out["Fe change (pts)"] == dl) & (out["Status"] == "Optimal")]
        if len(sub):
            r = sub.loc[sub["Cost Rs/tHM"].idxmin()]
            best.append({"Fe change (pts)": dl, "Best sinter %": r["Sinter %"], "Cost Rs/tHM": r["Cost Rs/tHM"], "Coke kg": r["Coke kg"]})
        else:
            best.append({"Fe change (pts)": dl, "Best sinter %": np.nan, "Cost Rs/tHM": np.nan, "Coke kg": np.nan})
    out.attrs["best"] = pd.DataFrame(best)
    out.attrs["group"] = group
    return out


def ore_price_scan(df, cfg, material, pcts=tuple(range(-40, 41, 10)), iters=12):
    """Price of one material moved across `pcts` (%) and the model re-solved each time.
    attrs['switches']: prices at which this material enters or leaves the burden (located by bisection)."""
    d0 = opt._fdf(opt.ensure_columns(df))
    if material not in d0.index or not bool(d0.loc[material, "Available"]):
        raise ValueError(f"'{material}' is not switched on.")
    p0 = float(d0.loc[material, "Price_Rs_t"])
    watch = on_materials(d0, "Iron_ore") + on_materials(d0, "Sinter")

    def at(pct):
        d = d0.copy()
        d.loc[material, "Price_Rs_t"] = max(1.0, p0 * (1 + pct / 100.0))
        return opt.solve(d, cfg, explain=False), float(d.loc[material, "Price_Rs_t"])

    rows, used = [], []
    for pc in sorted(float(x) for x in pcts):
        res, price = at(pc)
        r = {"Price change %": pc, "Price Rs/t": price, "Status": res[0]}
        if res[0] == "Optimal":
            r.update({"Cost Rs/tHM": res[2], "Sinter %": res[3]["Sinter_share_pct"], "Coke kg": res[3]["Coke_kg"]})
            for m in watch:
                r[f"{m} kg"] = res[1].get(m, 0.0)
            used.append((pc, res[1].get(material, 0.0) > 1e-3))
        rows.append(r)
    out = pd.DataFrame(rows)
    switches = []
    for (pa, ua), (pb, ub) in zip(used, used[1:]):
        if ua == ub:
            continue
        lo, hi = pa, pb
        for _ in range(iters):
            mid = (lo + hi) / 2
            res, _p = at(mid)
            u = res[0] == "Optimal" and res[1].get(material, 0.0) > 1e-3
            lo, hi = (mid, hi) if u == ua else (lo, mid)
        pct = (lo + hi) / 2
        switches.append({"Price change %": pct, "Price Rs/t": p0 * (1 + pct / 100.0),
                         "What happens": f"{material} {'leaves' if ua else 'enters'} the burden as the price rises past this point"})
    out.attrs["switches"] = switches
    out.attrs["base_price"] = p0
    return out


def marginal_effects(df, cfg, base_res):
    """The model's own 'rules of thumb' on this run's data: each input nudged once and the model re-solved under the
    same settings. Returns a table with the model's full effect next to what the thumb rule term alone gives."""
    if base_res[0] != "Optimal":
        raise ValueError("Needs an optimal run.")
    a0, c0 = base_res[3], base_res[2]
    d0 = opt.ensure_columns(df)
    rules = cfg.material_rules
    ref = cfg.ref_rates
    on = lambda k: bool(cfg.fuel_terms.get(k, True))
    rows = []

    def add(key, label, change, res, direct_fuel, rule_text, direct_cost=None):
        r = {"Key": key, "Effect of": label, "Change": change, "Status": res[0], "Thumb rule": rule_text,
             "Rule term alone, kg fuel": direct_fuel}
        if res[0] == "Optimal":
            a = res[3]
            r.update({"Coke kg": a["Coke_kg"] - a0["Coke_kg"], "Fuel kg": a["Fuel_supplied"] - a0["Fuel_supplied"],
                      "Cost Rs/tHM": res[2] - c0, "Sinter % after": a["Sinter_share_pct"]})
        else:
            r.update({"Coke kg": np.nan, "Fuel kg": np.nan, "Cost Rs/tHM": np.nan, "Sinter % after": np.nan})
        if direct_cost is not None:
            r["Rule term alone, Rs"] = direct_cost
        rows.append(r)

    for key, grp, col, label, rk in (("ore_fe", "Iron_ore", "Fe", "+1 point of ore Fe", "ore_fe"),
                                     ("sinter_fe", "Sinter", "Fe", "+1 point of sinter Fe", "sinter_fe"),
                                     ("ore_moisture", "Iron_ore", "Moisture_Pct", "+1 point of ore moisture", "ore_moisture"),
                                     ("coke_moisture", "Fuel_Coke", "Moisture_Pct", "+1 point of coke moisture", "coke_moisture")):
        if not on_materials(d0, grp):
            continue
        res = opt.solve(_shift(d0, grp, col, delta=1.0), cfg, explain=False)
        r = opt.MATERIAL_RULES[rk]
        coef = float(rules[rk]["coef"])
        pool_kg = {"Iron_ore": a0["Ore_kg"], "Sinter": a0["Sinter_kg"], "Fuel_Coke": a0["Coke_kg"] + a0["NutCoke_kg"]}[grp]
        direct = (r["sign"] * coef * pool_kg / ref[r["rate"]]) if on(rk) else 0.0
        txt = (f"{coef:g} kg fuel {'more' if r['sign'] > 0 else 'less'} per point at {ref[r['rate']]:g} kg/tHM"
               + ("" if on(rk) else " (rule switched off)"))
        add(key, label, f"all switched-on {grp.replace('_', ' ').replace('Fuel Coke', 'coke')} +1 pt", res, direct, txt)

    s0 = float(a0["Sinter_share_pct"])
    base_p = opt.solve(d0, pinned(cfg, s0), explain=False)
    done = False
    for s1 in (s0 + 10.0, s0 - 10.0):                   # up if the furnace can take it, otherwise down (sign reversed)
        if not (0.0 <= s1 <= 100.0) or base_p[0] != "Optimal":
            continue
        res1 = opt.solve(d0, pinned(cfg, s1), explain=False)
        if res1[0] != "Optimal":
            continue
        sign = 1.0 if s1 > s0 else -1.0
        a_b, a_1 = base_p[3], res1[3]
        rows.append({"Key": "sinter_share", "Effect of": "+10 points of sinter share",
                     "Change": f"sinter pinned at {min(s0, s1):.1f} % and {max(s0, s1):.1f} %", "Status": "Optimal",
                     "Coke kg": sign * (a_1["Coke_kg"] - a_b["Coke_kg"]), "Fuel kg": sign * (a_1["Fuel_supplied"] - a_b["Fuel_supplied"]),
                     "Cost Rs/tHM": sign * (res1[2] - base_p[2]), "Sinter % after": max(s0, s1),
                     "Thumb rule": f"{10 * cfg.fuel_per_pct_sinter:g} kg less fuel per +10 points" + ("" if on("sinter_share") else " (rule switched off)"),
                     "Rule term alone, kg fuel": -10.0 * cfg.fuel_per_pct_sinter if on("sinter_share") else 0.0})
        done = True
        break
    if not done:
        rows.append({"Key": "sinter_share", "Effect of": "+10 points of sinter share", "Change": "no feasible share 10 points either side",
                     "Status": "Infeasible", "Coke kg": np.nan, "Fuel kg": np.nan, "Cost Rs/tHM": np.nan, "Sinter % after": np.nan,
                     "Thumb rule": f"{10 * cfg.fuel_per_pct_sinter:g} kg less fuel per +10 points", "Rule term alone, kg fuel": np.nan})

    if on_materials(d0, "Fuel_Coke"):
        res = opt.solve(_shift(d0, "Fuel_Coke", "Price_Rs_t", add=1000.0), cfg, explain=False)
        add("coke_price", "+Rs 1,000/t coke price", "all switched-on coke +Rs 1,000/t", res, 0.0,
            "price only; the fuel rate does not depend on price", direct_cost=a0["Coke_kg"])
    out = pd.DataFrame(rows)
    order = ["ore_fe", "sinter_fe", "sinter_share", "ore_moisture", "coke_moisture", "coke_price"]
    out["_o"] = out["Key"].map({k: i for i, k in enumerate(order)})
    return out.sort_values("_o").drop(columns="_o").reset_index(drop=True)


def sinter_basicity_df(df, target_b2, dilute=True):
    """Every ON sinter set to basicity CaO/SiO2 = target_b2 by changing its CaO. dilute=True keeps the analysis a
    mass balance: the CaO added (or removed) per 100 kg of sinter dilutes (or concentrates) every other assay.
    Returns (new table, per-sinter summary)."""
    if target_b2 <= 0:
        raise ValueError("Target basicity must be above zero.")
    d = opt._fdf(opt.ensure_columns(df))
    sins = on_materials(d, "Sinter")
    if not sins:
        raise ValueError("No sinter is switched on.")
    rows = []
    for m in sins:
        cao, sio2, fe = float(d.loc[m, "CaO"]), float(d.loc[m, "SiO2"]), float(d.loc[m, "Fe"])
        if sio2 <= 0:
            raise ValueError(f"{m} has no SiO2, so its basicity cannot be set.")
        add = target_b2 * sio2 - cao                     # kg CaO added per 100 kg of the original sinter
        if dilute:
            if 100.0 + add <= 0:
                raise ValueError("That basicity would remove more CaO than the sinter holds.")
            k = 100.0 / (100.0 + add)
            for c_ in ("Fe", "MgO", "SiO2", "Al2O3", "Mn", "S", "FC"):
                d.loc[m, c_] = float(d.loc[m, c_]) * k
            d.loc[m, "CaO"] = (cao + add) * k
        else:
            d.loc[m, "CaO"] = max(0.0, cao + add)
        rows.append({"Sinter": m, "B2 now": cao / sio2, "B2 what-if": float(d.loc[m, "CaO"]) / float(d.loc[m, "SiO2"]),
                     "CaO % now": cao, "CaO % what-if": float(d.loc[m, "CaO"]), "Fe % now": fe, "Fe % what-if": float(d.loc[m, "Fe"]),
                     "SiO2 % now": sio2, "SiO2 % what-if": float(d.loc[m, "SiO2"])})
    return d, pd.DataFrame(rows)


# ================================================================ run history and saved scenarios
CFG_LABELS = {
    "hm_fe_pct": ("HM Fe %", 1), "hm_si_pct": ("HM Si %", 1), "hm_c_pct": ("HM C %", 1),
    "fe_required_per_100kg": ("Fe per 100 kg HM", 1), "fe_closure_basis": ("Fe closure basis", 1),
    "sinter_min": ("Sinter floor %", 100), "sinter_max": ("Sinter top %", 100), "sinter_manual_on": ("Sinter pin", 1),
    "sinter_manual_pct": ("Pinned sinter %", 1), "pci_fixed_kgthm": ("PCI kg", 1), "nut_coke_kgthm": ("Nut coke kg", 1),
    "nut_coke_mode": ("Nut coke rule", 1), "basicity_min": ("B2 min", 1), "basicity_max": ("B2 max", 1),
    "mgo_min_pct": ("MgO min %", 1), "mgo_max_pct": ("MgO max %", 1), "al2o3_min_pct": ("Al2O3 min %", 1),
    "al2o3_max_pct": ("Al2O3 max %", 1), "minor_max_kgthm": ("Minor cap kg", 1), "minor_demand_only": ("Minor on demand", 1),
    "stock_plan_hm_tonnes": ("Plan HM t", 1), "stock_balance": ("Stock balance", 1), "furnace_name": ("Furnace", 1),
    "slag_ref_kgthm": ("Slag reference kg", 1), "fuel_per_kg_slag": ("Slag rule kg/kg", 1), "sinter_ref_pct": ("Sinter reference %", 1),
    "fuel_per_pct_sinter": ("Sinter rule kg/pt", 1), "raw_flux_ref_kgthm": ("Raw flux reference kg", 1),
    "fuel_per_kg_raw_flux": ("Raw flux rule kg/kg", 1), "ks_fixed": ("Ks", 1), "price_basis": ("Price basis", 1), "om_rs_thm": ("O&M cost Rs/tHM", 1),
    "raw_flux_min_cao_mgo": ("Raw-flux threshold CaO+MgO %", 1), "dr_degree_floor": ("DRR for carbon floor", 1),
    "mn_reduction_eff": ("Mn reduction efficiency", 1), "mgo_al2o3_guide_lo": ("MgO/Al2O3 guide min", 1),
    "mgo_al2o3_guide_hi": ("MgO/Al2O3 guide max", 1), "fe_c_target": ("Fe/C guide", 1), "fe_c_tol": ("Fe/C guide tolerance", 1)}
HEAT_LABELS = {"mode": "Heat balance mode", **{k: v[0] for k, v in opt.HEAT_FIELDS.items()}}
DF_LABELS = {"Price_Rs_t": "price", "Moisture_Pct": "moisture", "RM_Stock": "stock", "Fines_Pct": "fines %",
             "Fines_Credit_Rs_t": "fines credit", "Fe": "Fe", "CaO": "CaO", "MgO": "MgO", "SiO2": "SiO2", "Al2O3": "Al2O3",
             "Mn": "Mn", "S": "S", "FC": "FC", "Group": "group"}


def _fmt(v):
    if isinstance(v, bool):
        return "on" if v else "off"
    if isinstance(v, (int, float, np.floating)):
        if isinstance(v, float) and math.isnan(v):
            return "blank"
        return f"{float(v):,.4g}" if abs(float(v)) < 1000 else f"{float(v):,.0f}"
    return str(v)


def describe_changes(old_df, new_df, old_cfg, new_cfg):
    """Plain list of what differs between two runs' inputs and settings (empty list = nothing)."""
    out = []
    if old_cfg is not None and new_cfg is not None:
        a, b = dataclasses.asdict(old_cfg), dataclasses.asdict(new_cfg)
        for k, (lab, scale) in CFG_LABELS.items():
            va, vb = a.get(k), b.get(k)
            if isinstance(va, float) or isinstance(vb, float):
                if va is not None and vb is not None and math.isclose(float(va), float(vb), abs_tol=1e-9):
                    continue
            elif va == vb:
                continue
            fa = va * scale if isinstance(va, (int, float)) and not isinstance(va, bool) else va
            fb = vb * scale if isinstance(vb, (int, float)) and not isinstance(vb, bool) else vb
            out.append(f"{lab} {_fmt(fa)} -> {_fmt(fb)}")
        for k, (lab, _desc) in opt.RULE_SWITCHES.items():
            if bool(a["fuel_terms"].get(k, True)) != bool(b["fuel_terms"].get(k, True)):
                out.append(f"{lab} rule {'on' if b['fuel_terms'].get(k, True) else 'off'}")
        for k in b["material_rules"]:
            for part in ("coef", "ref"):
                va, vb = a["material_rules"].get(k, {}).get(part), b["material_rules"][k][part]
                if va is None or not math.isclose(float(va), float(vb), abs_tol=1e-9):
                    out.append(f"{opt.TERM_LABELS.get(k, k)} {'kg/pt' if part == 'coef' else 'reference'} {_fmt(va)} -> {_fmt(vb)}")
        if a["furnace_profiles"] != b["furnace_profiles"]:
            for fn in sorted(set(a["furnace_profiles"]) | set(b["furnace_profiles"])):
                va, vb = a["furnace_profiles"].get(fn, {}).get("base"), b["furnace_profiles"].get(fn, {}).get("base")
                if va != vb:
                    out.append(f"{fn} base fuel {_fmt(va) if va is not None else 'none'} -> {_fmt(vb) if vb is not None else 'removed'}")
        for k in sorted(set(a.get("ref_rates", {})) | set(b.get("ref_rates", {}))):
            va, vb = a.get("ref_rates", {}).get(k), b.get("ref_rates", {}).get(k)
            if va is None or vb is None or not math.isclose(float(va), float(vb), abs_tol=1e-9):
                out.append(f"reference {k} rate {_fmt(va)} -> {_fmt(vb)} kg/tHM")
        ha, hb = a.get("heat", {}), b.get("heat", {})
        for k, lab in HEAT_LABELS.items():
            va, vb = ha.get(k), hb.get(k)
            if isinstance(va, (int, float)) and isinstance(vb, (int, float)) and not isinstance(va, bool):
                if math.isclose(float(va), float(vb), abs_tol=1e-9):
                    continue
            elif va == vb:
                continue
            out.append(f"{lab} {_fmt(va)} -> {_fmt(vb)}")
    if old_df is not None and new_df is not None:
        o, n = opt.ensure_columns(old_df), opt.ensure_columns(new_df)
        for m in n.index.difference(o.index):
            out.append(f"{m} added")
        for m in o.index.difference(n.index):
            out.append(f"{m} removed")
        for m in n.index.intersection(o.index):
            if bool(o.loc[m, "Available"]) != bool(n.loc[m, "Available"]):
                out.append(f"{m} switched {'on' if bool(n.loc[m, 'Available']) else 'off'}")
            for c_, lab in DF_LABELS.items():
                va, vb = o.loc[m, c_], n.loc[m, c_]
                if c_ == "Group":
                    if str(va) != str(vb):
                        out.append(f"{m} {lab} {va} -> {vb}")
                    continue
                if pd.isna(va) and pd.isna(vb):
                    continue
                if pd.isna(va) or pd.isna(vb) or not math.isclose(float(va), float(vb), abs_tol=1e-9):
                    out.append(f"{m} {lab} {_fmt(va)} -> {_fmt(vb)}")
    return out


def change_label(changes, limit=3):
    if not changes:
        return "no change"
    more = len(changes) - limit
    return "; ".join(changes[:limit]) + (f" (+{more} more)" if more > 0 else "")


def scenario_from_bundle(name, bundle):
    """Freeze an optimal run as a named scenario (inputs, settings, key numbers, recipe and 0-100 % curve)."""
    if bundle is None or bundle.get("status") != "Optimal":
        raise ValueError("Only an optimal run can be saved as a scenario.")
    a = bundle["ach"]
    bf = bundle["burden_full"]
    recipe = bf[(bf["Result"] == "USED")].set_index("Material")["Dry kg/tHM"].astype(float).to_dict()
    kpis = {"Cost Rs/tHM": bundle["cost"], "Sinter %": a["Sinter_share_pct"], "Regular coke kg": a["Coke_kg"],
            "Total fuel kg": a["Fuel_supplied"], "Slag kg": a["slag_kg"], "B2": a["B2"], "MgO %": a["MgO_pct"],
            "Al2O3 %": a["Al2O3_pct"], "Raw flux kg": a["Raw_flux_kg"], "Flux kg": a["Flux_kg"], "Ore kg": a["Ore_kg"],
            "Sinter kg": a["Sinter_kg"]}
    return {"name": str(name).strip() or "Scenario", "created": bundle["created"], "kpis": kpis, "recipe": recipe,
            "groups": bf.set_index("Material")["Group"].to_dict(), "cfg": copy.deepcopy(bundle["cfg"]),
            "inputs": bundle["inputs"].copy(), "curve": bundle.get("curve"), "furnace": bundle["furnace"]}


def compare_scenarios(scens):
    """(key numbers table, recipe table, settings differences) for 1-3 scenarios; columns are scenario names."""
    names = [s["name"] for s in scens]
    kp = pd.DataFrame({s["name"]: s["kpis"] for s in scens})
    kp.index.name = "Result"
    kp = kp.reset_index()
    if len(scens) >= 2:
        first = names[0]
        for n in names[1:]:
            kp[f"{n} vs {first}"] = kp[n] - kp[first]
    mats = []
    for s in scens:
        for m in s["recipe"]:
            if m not in mats:
                mats.append(m)
    order = {g: i for i, g in enumerate(opt.GROUPS)}
    groups = {}
    for s in scens:
        groups.update(s["groups"])
    mats.sort(key=lambda m: (order.get(str(groups.get(m)), 99), m))
    rc = pd.DataFrame([{"Material": m, "Group": groups.get(m, ""), **{s["name"]: s["recipe"].get(m, 0.0) for s in scens}} for m in mats])
    diffs = []
    for s in scens[1:]:
        ch = describe_changes(scens[0]["inputs"], s["inputs"], scens[0]["cfg"], s["cfg"])
        diffs.append((s["name"], scens[0]["name"], ch))
    return kp, rc, diffs


# ================================================================ extra Excel sheets
def extra_sheets(items):
    """A callable for optimiser.export_bytes(..., extra=...) that adds one formatted sheet per (name, title, df, note)."""
    items = [it for it in items if it and it[2] is not None and len(it[2])]

    def add(wb):
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter as L
        navy, light, pale, grid = "1F3864", "DDEBF7", "F3F8FD", "BFBFBF"
        fill = lambda h: PatternFill("solid", start_color=h, end_color=h)
        thin = Side(style="thin", color=grid)
        box = Border(left=thin, right=thin, top=thin, bottom=thin)
        for name, title, df, note in items:
            ws = wb.create_sheet(name[:31])
            n = max(len(df.columns), 2)
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
            c = ws.cell(1, 1, title)
            c.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
            c.alignment = Alignment(vertical="center", indent=1)
            ws.row_dimensions[1].height = 26
            for j in range(1, n + 1):
                ws.cell(1, j).fill = fill(navy)
            if note:
                ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
                c = ws.cell(2, 1, note)
                c.font = Font(name="Calibri", size=10, italic=True, color="44546A")
                c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
                ws.row_dimensions[2].height = 30
                for j in range(1, n + 1):
                    ws.cell(2, j).fill = fill(light)
            r0 = 4
            widths = {}
            for j, col in enumerate(df.columns, start=1):
                c = ws.cell(r0, j, str(col))
                c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
                c.fill = fill(navy)
                c.border = box
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                widths[j] = min(len(str(col)), 24)
            ws.row_dimensions[r0].height = 30
            for i, row in enumerate(df.itertuples(index=False), start=1):
                for j, v in enumerate(row, start=1):
                    if isinstance(v, (np.floating, float)):
                        v = None if (math.isnan(float(v)) or math.isinf(float(v))) else float(v)
                    elif isinstance(v, np.integer):
                        v = int(v)
                    elif isinstance(v, (bool, np.bool_)):
                        v = "Yes" if v else "No"
                    elif v is not None and not isinstance(v, (int, str)):
                        v = str(v)
                    c = ws.cell(r0 + i, j, v)
                    c.border = box
                    c.font = Font(name="Calibri", size=10, bold=str(row[0]).startswith("TOTAL"))
                    if isinstance(v, float):
                        c.number_format = "#,##0.00"
                    if i % 2 == 0:
                        c.fill = fill(pale)
                    if v is not None:
                        widths[j] = max(widths[j], len(f"{v:,.2f}") if isinstance(v, float) else len(str(v)))
            for j, w in widths.items():
                ws.column_dimensions[L(j)].width = max(10, min(48, w + 2.5))
            ws.freeze_panes = ws.cell(r0 + 1, 2)
            ws.sheet_view.showGridLines = False
            ws.sheet_properties.tabColor = "7030A0"
    return add
