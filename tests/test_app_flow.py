"""The whole app, driven the way a user would (Streamlit AppTest), on the real input files."""
import io
import time

import pytest
from streamlit.testing.v1 import AppTest

from helpers import ROOT, SINTER_ON, real_inputs
from combined import loop as L

PAGES = {
    "combined": ["Uploads & shared settings", "Optimise hot metal cost", "Hot metal cost", "Basicity scan", "Price drivers", "Saved scenarios", "Export", "Glossary"],
    "impact": ["Baseline & changes", "Impact", "Optimise sinter inputs", "What each condition costs", "Impact scenarios", "Export"],
    "sinter": ["Upload & Settings", "Inputs", "RM Stock & Materials", "Dashboard", "Recipe & Composition", "Inventory Usage", "Manual Burden Control",
               "Scenario Analysis", "Plant Run Validation", "Productivity", "Wet Specific Consumption", "Reports"],
    "mbf": ["Upload & settings", "Inputs", "Materials & stock", "Dashboard", "Burden & cost", "Slag oxides", "Trends", "Scenario analysis",
            "Reports & export", "Heat audit"],
}


def fresh_app():
    at = AppTest.from_file(f"{ROOT}/app.py", default_timeout=300)
    at.run()
    assert not at.exception
    sdf, mdf = real_inputs()
    at.session_state["sinter__master_df"] = sdf
    at.session_state["sinter__source"] = "SInter_Input.xlsx"
    at.session_state["sinter__available"] = {x: (x in SINTER_ON) for x in sdf.index}
    at.session_state["mbf__df"] = mdf
    at.session_state["mbf__source"] = "MBF_Input.xlsx"
    at.session_state["mbf__demo"] = False
    at.session_state["mbf__mat_ver"] = at.session_state["mbf__mat_ver"] + 1
    return at


def go(at, ws, page):
    at.session_state["ws"] = ws
    at.session_state[f"route_{ws}"] = page
    at.run()


def press(at, label):
    [b for b in at.button if b.label == label][0].click()
    at.run()


def text(at):
    return " ".join(x.value for x in at.markdown)


@pytest.fixture(scope="module")
def app():
    return fresh_app()


@pytest.mark.parametrize("ws,page", [(w, p) for w, ps in PAGES.items() for p in ps])
def test_every_page_renders(app, ws, page):
    go(app, ws, page)
    assert not app.exception, app.exception[0].value if app.exception else ""


def test_state_of_the_two_models_is_separate(app):
    assert "sinter__result" in app.session_state and "mbf__result" in app.session_state
    assert "result" not in app.session_state and "df" not in app.session_state          # nothing unprefixed leaks


def test_original_pages_keep_all_their_widgets(app):
    go(app, "sinter", "Inputs")
    assert len(app.number_input) == 37 and len(app.tabs) == 6
    go(app, "mbf", "Inputs")
    assert len(app.number_input) == 53 and len(app.tabs) == 5       # 51 + the furnace O&M input + the O&M box at the top of every MBF page


def test_mbf_cannot_run_before_the_sinter_model(app):
    go(app, "mbf", "Dashboard")
    t = text(app)
    assert "has not been run" in t and "Fix these before running" in t


def test_sinter_result_replaces_the_mbf_sinter_row():
    at = fresh_app()
    go(at, "sinter", "Dashboard")
    [b for b in at.button if "RUN OPTIMIZER" in b.label.upper()][0].click()
    at.run()
    assert at.session_state["sinter__runs"] == 1
    go(at, "mbf", "Materials & stock")
    df = at.session_state["mbf__df"]
    row = L.find_sinter_row(df)
    res = at.session_state["sinter__result"]
    assert df.loc[row, "Price_Rs_t"] == pytest.approx(res["cost"] + 750.0, abs=0.01)
    assert df.loc[row, "Fe"] == pytest.approx(res["achieved"]["Fe"], abs=1e-6)
    assert df.loc[row, "Moisture_Pct"] == 3.0                                            # stays from the MBF Excel
    go(at, "mbf", "Dashboard")
    press(at, "Run optimiser")
    assert not at.exception and at.session_state["mbf__result"][0] == "Optimal"
    assert at.session_state["mbf__result"][2] == pytest.approx(31038.7, abs=1.0)         # the Combined screenshot's furnace figure
    at.session_state["mbf_sinter_mode"] = "Typed row (original behaviour)"
    at.run()
    df = at.session_state["mbf__df"]
    assert float(df.loc[row, "Price_Rs_t"]) == 7500.0 and float(df.loc[row, "Fe"]) == pytest.approx(52.57)   # typed values restored


def test_run_both_models_then_sinter_price_moves_the_cost():
    at = fresh_app()
    go(at, "combined", "Hot metal cost")
    assert "No run yet" in text(at)
    press(at, "Run both models")
    assert not at.exception
    run = at.session_state["cmb_run"]
    sm = L.summary(run)
    assert run["ok"] and run["converged"]
    t = text(at)
    for needle in ("Combined result", "Where the money goes", "Alerts", "Quality limits", "Mill scale used", "Fe achieved", "Coke used", "Iron ore used",
                   "Sinter in the sinter + ore mix", "PCI used", "Slag volume", "Flux used", "Nut coke used", "fuel rate"):
        assert needle in t
    assert "Iron ore used (mill scale included)" not in t                                  # sinter panel no longer shows the iron-ore block
    assert "Stock check" not in t                                                         # removed from this page (the Export workbook still has the two stock-check sheets)
    # the acceptance test: change a sinter input, run, and the cost per tonne of hot metal moves
    sdf = at.session_state["sinter__master_df"].copy()
    sdf["Price_Rs_t"] = sdf["Price_Rs_t"].astype(float)
    sdf.loc["KIOM", "Price_Rs_t"] *= 1.10
    at.session_state["sinter__master_df"] = sdf
    at.run()
    assert "Inputs changed since this run" in text(at)                                    # stale detection
    press(at, "Run both models")
    sm2 = L.summary(at.session_state["cmb_run"])
    assert sm2["hm_cost"] > sm["hm_cost"] + 100 and sm2["sinter_price"] > sm["sinter_price"] + 100
    assert "KIOM price" in text(at)                                                       # the change panel names the cause
    # management view keeps the KPI cards and quality limits and hides the engineer detail
    at.session_state["cmb_view"] = "Management view"
    at.run()
    assert "Where the money goes" in text(at) and "Quality limits" in text(at) and "Cost breakdown" not in text(at)
    # scenarios and export
    go(at, "combined", "Saved scenarios")
    press(at, "Save this run")
    assert len(at.session_state["cmb_scen"]) == 1
    go(at, "combined", "Export")
    press(at, "Export combined results")
    assert not at.exception
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(at.session_state["cmb_export"][0]))
    assert {"Combined Summary", "Sinter Recipe", "Loop Passes", "Stock Check Sinter", "Stock Check Furnace"} <= set(wb.sheetnames)


def test_optimise_page_works_from_the_combined_run():
    at = fresh_app()
    go(at, "combined", "Hot metal cost")
    press(at, "Run both models")
    # the optimise page works from this run (it says so honestly when nothing is left to gain)
    go(at, "impact", "Optimise sinter inputs")
    assert "Find the lowest hot metal cost" in [b.label for b in at.button]
    press(at, "Find the lowest hot metal cost")
    assert not at.exception, at.exception[0].value if at.exception else ""
    opt = at.session_state["cmb_opt"]
    assert opt["ok"] and opt["base"]["ok"]
    t = text(at)
    assert "What the furnace pays for sinter chemistry" in t and "current recipe stands" in t.replace("\n", " ")


def _tag(at):
    run = at.session_state["cmb_run"]
    return f"{run['time']:%H%M%S%f}_{at.session_state['imp_reset']}"


def test_impact_model_end_to_end():
    at = fresh_app()
    go(at, "combined", "Hot metal cost")
    press(at, "Run both models")
    assert at.session_state["cmb_run"]["ok"]
    base_run = at.session_state["cmb_run"]
    hm0 = L.summary(base_run)["hm_cost"]
    # no run yet: honest empty states
    go(at, "impact", "Impact")
    assert "No impact run yet" in text(at)
    # change the sinter O&M and run: lands on the Impact page with the bridge
    go(at, "impact", "Baseline & changes")
    assert not at.exception and "Run impact" in [b.label for b in at.button]
    [n for n in at.number_input if n.key == f"imp_om_{_tag(at)}"][0].set_value(900.0)
    at.run()
    assert "Sinter O&M" in text(at) or any("Sinter O&M" in str(df.value) for df in at.dataframe)
    press(at, "Run impact")
    assert not at.exception, at.exception[0].value if at.exception else ""
    res = at.session_state["imp_res"]
    assert res["ok"] and res["total"] > 100.0                                      # 150 Rs/t more sinter O&M, about 1.3 t of sinter per tHM
    assert at.session_state["route_impact"] == "Impact"
    t = text(at)
    for needle in ("Which change did it", "How it travels from sinter making to the furnace", "What the furnace did about it"):
        assert needle in t
    assert at.session_state["cmb_run"] is base_run                                 # the combined run was not touched
    assert L.summary(at.session_state["cmb_run"])["hm_cost"] == hm0
    # the other pages of the model work from it
    go(at, "impact", "Impact scenarios")
    press(at, "Save this result")
    assert len(at.session_state["imp_scen"]) == 1
    go(at, "impact", "What each condition costs")
    press(at, "Measure what each condition costs")
    assert not at.exception and len(at.session_state["imp_sweep"]["tbl"]) >= 5
    go(at, "impact", "Export")
    press(at, "Prepare workbook")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(at.session_state["imp_export"][0]))
    assert {"Summary", "Changes", "By path", "Condition prices", "Saved results"} <= set(wb.sheetnames)



def test_furnace_om_is_editable_at_the_top_of_every_mbf_page():
    at = fresh_app()
    for page in ("Dashboard", "Inputs", "Burden & cost"):
        go(at, "mbf", page)
        assert any(n.label == "Furnace O&M, Rs per tHM" for n in at.number_input), page
    go(at, "mbf", "Dashboard")
    [n for n in at.number_input if n.label == "Furnace O&M, Rs per tHM"][0].set_value(1200.0)
    at.run()
    assert not at.exception and at.session_state["mbf__cfg"].om_rs_thm == 1200.0
    go(at, "mbf", "Inputs")                                                       # the same setting, shown on the page where it always was
    assert any(abs(n.value - 1200.0) < 1e-9 for n in at.number_input if "O&M" in n.label)
    go(at, "combined", "Uploads & shared settings")
    assert any(n.label == "Furnace O&M, Rs per tHM" and abs(n.value - 1200.0) < 1e-9 for n in at.number_input)
